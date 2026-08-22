import calendar
from io import BytesIO
from datetime import date, datetime

from flask import request, send_file

from ...extensions import db
from ...models import AuditLog, ImportBatch, ImportError, Teacher, TeacherDetail, TeacherTier
from ...services.teacher_service import create_teacher as svc_create_teacher, generate_teacher_no
from ...utils.storage import file_url as _file_url, storage_reference
from .helpers import current_admin_id, require_admin_roles, require_admin_token, _date_text
from . import admin_bp


@admin_bp.get("/teachers")
@require_admin_token
def teacher_list():
    teachers = Teacher.query.filter(Teacher.status != "hidden").order_by(Teacher.teacher_no.asc()).all()
    items = [
        {
            "id": t.id,
            "teacherNo": t.teacher_no,
            "name": t.real_name,
            "xileName": t.xile_name,
            "tier": t.tier.code if t.tier else None,
            "tierName": t.tier.name if t.tier else None,
            "city": t.city,
            "district": t.district,
            "status": t.status,
            "validUntil": _date_text(t.valid_until),
            "certifiedAt": _date_text(t.first_certified_on),
            "avatarUrl": storage_reference(t.avatar_url),
            "certificateUrl": _file_url(t.certificate_url),
            "phone": t.detail.phone if t.detail else None,
            "committeeRemark": t.detail.committee_remark if t.detail else None,
        }
        for t in teachers
    ]
    return {"items": items, "total": len(items)}


@admin_bp.get("/teachers/<int:teacher_id>")
@require_admin_token
def get_teacher_detail(teacher_id):
    teacher = db.session.get(Teacher, teacher_id)
    if teacher is None or teacher.status == "hidden":
        return {"error": "not found"}, 404
    return {
        "id": teacher.id,
        "teacherNo": teacher.teacher_no,
        "name": teacher.real_name,
        "xileName": teacher.xile_name,
        "tier": teacher.tier.code if teacher.tier else None,
        "tierName": teacher.tier.name if teacher.tier else None,
        "city": teacher.city,
        "district": teacher.district,
        "status": teacher.status,
        "validUntil": _date_text(teacher.valid_until),
        "certifiedAt": _date_text(teacher.first_certified_on),
        "avatarUrl": storage_reference(teacher.avatar_url),
        "certificateUrl": _file_url(teacher.certificate_url),
        "phone": teacher.detail.phone if teacher.detail else None,
        "specialties": teacher.detail.specialties if teacher.detail else None,
        "teachingSummary": teacher.detail.teaching_summary if teacher.detail else None,
        "committeeRemark": teacher.detail.committee_remark if teacher.detail else None,
    }


@admin_bp.put("/teachers/<int:teacher_id>")
@require_admin_token
@require_admin_roles("admin", "super_admin")
def update_teacher(teacher_id):
    teacher = db.session.get(Teacher, teacher_id)
    if teacher is None or teacher.status == "hidden":
        return {"error": "not found"}, 404

    payload = request.get_json(silent=True) or {}

    if "name" in payload:
        teacher.real_name = str(payload["name"] or "").strip()
    if "xileName" in payload:
        teacher.xile_name = str(payload["xileName"] or "").strip() or None
    if "city" in payload:
        teacher.city = str(payload["city"] or "").strip() or None
    if "district" in payload:
        teacher.district = str(payload["district"] or "").strip() or None
    if "avatarUrl" in payload:
        teacher.avatar_url = storage_reference((payload["avatarUrl"] or "").strip())
    if "certificateUrl" in payload:
        teacher.certificate_url = storage_reference((payload["certificateUrl"] or "").strip())

    if "committeeRemark" in payload or "phone" in payload or "specialties" in payload:
        if not teacher.detail:
            detail = TeacherDetail(teacher_id=teacher.id)
            db.session.add(detail)
            db.session.flush()
        if "committeeRemark" in payload:
            teacher.detail.committee_remark = str(payload["committeeRemark"] or "").strip() or None
        if "phone" in payload:
            teacher.detail.phone = str(payload["phone"] or "").strip() or None
        if "specialties" in payload:
            teacher.detail.specialties = str(payload["specialties"] or "").strip() or None

    db.session.commit()
    db.session.add(AuditLog(admin_id=current_admin_id() or 1, action="update_teacher", target_type="teacher", target_id=teacher.id))
    db.session.commit()
    return {"id": teacher.id, "name": teacher.real_name}


@admin_bp.post("/teachers")
@require_admin_token
@require_admin_roles("admin", "super_admin")
def create_teacher():
    payload = request.get_json(silent=True) or {}
    name = (payload.get("name") or "").strip()
    if not name:
        return {"error": "name required"}, 400

    valid_until_str = payload.get("expiryDate")
    valid_until = None
    if valid_until_str:
        try:
            valid_until = date.fromisoformat(valid_until_str.replace(".", "-"))
        except ValueError:
            pass

    teacher, tier = svc_create_teacher(
        name=name,
        tier_code=payload.get("level", "L1"),
        city=payload.get("city", "").strip(),
        district=payload.get("district", "").strip(),
        xile_name=payload.get("xileName", "").strip(),
        phone=payload.get("phone", "").strip(),
        valid_until=valid_until,
    )

    db.session.add(AuditLog(admin_id=1, action="create_teacher", target_type="teacher", target_id=teacher.id))
    db.session.commit()

    return {
        "id": teacher.id,
        "teacherNo": teacher.teacher_no,
        "name": teacher.real_name,
        "tier": tier.code,
        "validUntil": _date_text(teacher.valid_until),
    }, 201


@admin_bp.delete("/teachers/<int:teacher_id>")
@require_admin_token
@require_admin_roles("admin", "super_admin")
def delete_teacher(teacher_id):
    teacher = db.session.get(Teacher, teacher_id)
    if teacher is None:
        return {"error": "not found"}, 404
    teacher.status = "hidden"
    db.session.add(AuditLog(admin_id=1, action="delete_teacher", target_type="teacher", target_id=teacher.id))
    db.session.commit()
    return {"id": teacher.id, "status": "hidden"}


# ─── Import ──────────────────────────────────────────────────────────────────

VALID_TIERS = {"L0", "L1", "L2", "L3", "L4", "L5"}

# Expected column order:
# A: name, B: phone, C: tier, D: city, E: district,
# F: certifiedAt, G: validUntil, H: xileName, I: teacherNo


def _cell_str(row, idx):
    """Safely extract a stripped string from a row cell, handling short rows."""
    if idx >= len(row) or row[idx] is None:
        return ""
    val = row[idx]
    if isinstance(val, datetime):
        return val.strftime("%Y-%m-%d")
    if isinstance(val, date):
        return val.isoformat()
    if isinstance(val, float) and val == int(val):
        return str(int(val))
    return str(val).strip()


def _parse_date(text):
    """Parse a date string accepting YYYY-MM-DD, YYYY.MM.DD, YYYY/MM/DD formats."""
    if not text:
        return None
    normalized = text.replace(".", "-").replace("/", "-")
    try:
        return date.fromisoformat(normalized)
    except ValueError:
        return None


def _is_empty_row(row):
    """Return True if all cells in the row are None or whitespace-only."""
    if row is None:
        return True
    return all(cell is None or str(cell).strip() == "" for cell in row)


def _parse_import_row(row):
    """Parse a single Excel row into a dict of field values."""
    return {
        "name": _cell_str(row, 0),
        "phone": _cell_str(row, 1),
        "tier": _cell_str(row, 2).upper(),
        "city": _cell_str(row, 3),
        "district": _cell_str(row, 4),
        "certifiedAt": _cell_str(row, 5),
        "validUntil": _cell_str(row, 6),
        "xileName": _cell_str(row, 7),
        "teacherNo": _cell_str(row, 8),
    }


@admin_bp.get("/import/teachers/template")
@require_admin_token
def import_template():
    """Download the canonical teacher import workbook."""
    try:
        import openpyxl
        from openpyxl.styles import Font
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "教师资料"
        headers = ["姓名", "手机号", "认证等级", "城市", "地区", "首次认证日期", "有效期至", "喜乐名", "教师编号"]
        sheet.append(headers)
        sheet.append(["张三", "13800000000", "L1", "上海", "浦东新区", "2024-01-01", "2026-12-31", "张三老师", ""])
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for index, width in enumerate([16, 16, 12, 14, 16, 18, 16, 16, 16], start=1):
            sheet.column_dimensions[openpyxl.utils.get_column_letter(index)].width = width
        output = BytesIO()
        workbook.save(output)
        output.seek(0)
        return send_file(output, as_attachment=True, download_name="教师批量导入模板.xlsx", mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    except Exception:
        return {"error": "导入模板生成失败"}, 500


@admin_bp.post("/import/teachers/preview")
@require_admin_token
@require_admin_roles("admin", "super_admin")
def import_preview():
    file = request.files.get("file")
    if not file:
        return {"error": "file required"}, 400

    try:
        import openpyxl
        wb = openpyxl.load_workbook(file, read_only=True, data_only=True)
        ws = wb.active
    except Exception:
        return {"error": "invalid excel file"}, 400

    rows = list(ws.iter_rows(min_row=2, values_only=True))
    wb.close()

    errors = []
    parsed_rows = []
    seen_teacher_nos = set()

    # Collect existing teacherNo values from DB for duplicate detection
    existing_nos = set(
        no for (no,) in db.session.query(Teacher.teacher_no).all()
    )

    for idx, row in enumerate(rows, start=2):
        # Skip completely empty rows
        if _is_empty_row(row):
            continue

        data = _parse_import_row(row)
        row_errors = []

        # Required: name
        if not data["name"]:
            row_errors.append({"rowNumber": idx, "field": "name", "message": "姓名不能为空"})

        # Required: tier (must be L0-L5)
        if not data["tier"] or data["tier"] not in VALID_TIERS:
            row_errors.append({"rowNumber": idx, "field": "tier", "message": "等级代码无效（需为L0-L5）"})

        # Required: certifiedAt
        cert_date = _parse_date(data["certifiedAt"])
        if not data["certifiedAt"]:
            row_errors.append({"rowNumber": idx, "field": "certifiedAt", "message": "认证日期不能为空"})
        elif cert_date is None:
            row_errors.append({"rowNumber": idx, "field": "certifiedAt", "message": "日期格式无效（需YYYY-MM-DD）"})

        # Optional: validUntil (validate format if provided)
        if data["validUntil"]:
            valid_until = _parse_date(data["validUntil"])
            if valid_until is None:
                row_errors.append({"rowNumber": idx, "field": "validUntil", "message": "有效期格式无效（需YYYY-MM-DD）"})

        # Optional: teacherNo (check duplicates)
        if data["teacherNo"]:
            if data["teacherNo"] in seen_teacher_nos:
                row_errors.append({"rowNumber": idx, "field": "teacherNo", "message": "编号在文件中重复"})
            elif data["teacherNo"] in existing_nos:
                row_errors.append({"rowNumber": idx, "field": "teacherNo", "message": "编号已存在于系统中"})
            else:
                seen_teacher_nos.add(data["teacherNo"])

        if row_errors:
            errors.extend(row_errors)
        else:
            parsed_rows.append(data)

    total_rows = len([r for r in rows if not _is_empty_row(r)])
    valid_rows = len(parsed_rows)

    batch = ImportBatch(
        admin_id=1,
        total_rows=total_rows,
        success_rows=valid_rows,
        error_rows=total_rows - valid_rows,
        status="preview",
    )
    db.session.add(batch)
    db.session.flush()

    for err in errors:
        db.session.add(ImportError(
            batch_id=batch.id,
            row_number=err["rowNumber"],
            field=err["field"],
            error_message=err["message"],
        ))

    db.session.commit()

    return {
        "batchId": batch.id,
        "totalRows": total_rows,
        "validRows": valid_rows,
        "errors": errors,
        "preview": parsed_rows[:50],  # Return first 50 rows for UI preview
    }


@admin_bp.post("/import/teachers/commit")
@require_admin_token
@require_admin_roles("admin", "super_admin")
def import_commit():
    payload = request.get_json(silent=True) or {}
    batch_id = payload.get("batchId") or request.form.get("batchId")
    if batch_id is None:
        return {"error": "batchId required"}, 400

    batch = db.session.get(ImportBatch, batch_id)
    if batch is None or batch.status != "preview":
        return {"error": "invalid batch"}, 400

    file = request.files.get("file")
    if not file:
        return {"error": "file required for commit"}, 400

    try:
        import openpyxl
        wb = openpyxl.load_workbook(file, read_only=True, data_only=True)
        ws = wb.active
    except Exception:
        return {"error": "invalid excel file"}, 400

    rows = list(ws.iter_rows(min_row=2, values_only=True))
    wb.close()

    # Collect existing teacherNo values for duplicate check during commit
    existing_nos = set(
        no for (no,) in db.session.query(Teacher.teacher_no).all()
    )
    used_nos = set()
    created = 0

    for row in rows:
        # Skip empty rows
        if _is_empty_row(row):
            continue

        data = _parse_import_row(row)

        # Validate required fields — skip invalid rows silently
        if not data["name"]:
            continue
        if not data["tier"] or data["tier"] not in VALID_TIERS:
            continue

        cert_date = _parse_date(data["certifiedAt"])
        if cert_date is None:
            continue

        tier = TeacherTier.query.filter_by(code=data["tier"]).first()
        if not tier:
            continue

        # Determine teacherNo: use provided value or generate
        teacher_no = data["teacherNo"]
        if teacher_no:
            # Skip if duplicate
            if teacher_no in existing_nos or teacher_no in used_nos:
                continue
            used_nos.add(teacher_no)
        else:
            teacher_no = generate_teacher_no()
            # Ensure generated number is tracked to avoid collision within batch
            existing_nos.add(teacher_no)

        # Determine validUntil: use provided date or calculate from tier cycle
        valid_until = _parse_date(data["validUntil"])
        if valid_until is None:
            cycle = tier.review_cycle_years or 3
            target_year = cert_date.year + cycle
            target_day = min(cert_date.day, calendar.monthrange(target_year, cert_date.month)[1])
            valid_until = date(target_year, cert_date.month, target_day)

        teacher = Teacher(
            teacher_no=teacher_no,
            real_name=data["name"],
            xile_name=data["xileName"] or None,
            tier_id=tier.id,
            city=data["city"] or None,
            district=data["district"] or None,
            status="active",
            first_certified_on=cert_date,
            valid_until=valid_until,
        )
        db.session.add(teacher)
        db.session.flush()

        if data["phone"]:
            db.session.add(TeacherDetail(teacher_id=teacher.id, phone=data["phone"]))

        created += 1

    batch.status = "committed"
    batch.success_rows = created
    db.session.add(AuditLog(admin_id=1, action="import_teachers", target_type="batch", target_id=batch.id))
    db.session.commit()

    return {"batchId": batch.id, "createdCount": created}
