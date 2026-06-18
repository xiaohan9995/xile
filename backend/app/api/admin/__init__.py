from functools import wraps
from datetime import date, datetime, timedelta, timezone

from flask import Blueprint, current_app, request

from ...extensions import db
from ...models import (
    AdminUser,
    AnnualReview,
    AuditLog,
    ImportBatch,
    ImportError,
    Studio,
    SystemConfig,
    Teacher,
    TeacherDetail,
    TeacherTier,
)

admin_bp = Blueprint("admin", __name__)


def require_admin_token(view):
    @wraps(view)
    def wrapped(*args, **kwargs):
        expected_token = current_app.config.get("ADMIN_DEV_TOKEN")
        auth_header = request.headers.get("Authorization", "")
        if not expected_token or auth_header != f"Bearer {expected_token}":
            return {"error": "unauthorized"}, 401
        return view(*args, **kwargs)

    return wrapped


def _date_text(value):
    return value.strftime("%Y.%m.%d") if value else None


def _datetime_text(value):
    return value.strftime("%Y.%m.%d %H:%M") if value else None


# ─── Auth ────────────────────────────────────────────────────────────────────


@admin_bp.post("/login")
def login():
    payload = request.get_json(silent=True) or {}
    username = payload.get("username", "")
    password = payload.get("password", "")

    if username != "admin" or password != "password":
        return {"error": "invalid credentials"}, 401

    token = current_app.config.get("ADMIN_DEV_TOKEN") or "dev-admin-token"
    return {
        "token": token,
        "admin": {"username": "admin", "name": "系统管理员"},
    }


# ─── Dashboard ───────────────────────────────────────────────────────────────


@admin_bp.get("/stats/dashboard")
@require_admin_token
def dashboard_stats():
    today = date.today()
    expiring_deadline = today + timedelta(days=90)
    current_year = today.year

    return {
        "teacherCount": Teacher.query.filter(Teacher.status != "hidden").count(),
        "activeTeacherCount": Teacher.query.filter_by(status="active").count(),
        "pendingReviewCount": AnnualReview.query.filter_by(status="submitted").count(),
        "openStudioCount": Studio.query.filter_by(status="open").count(),
        "expiringCount": Teacher.query.filter(
            Teacher.status == "active",
            Teacher.valid_until <= expiring_deadline,
            Teacher.valid_until >= today,
        ).count(),
        "completedReviewCount": AnnualReview.query.filter(
            AnnualReview.status == "approved",
            db.extract("year", AnnualReview.reviewed_at) == current_year,
        ).count(),
    }


# ─── Teachers ────────────────────────────────────────────────────────────────


@admin_bp.get("/teachers")
@require_admin_token
def teacher_list():
    teachers = Teacher.query.filter(Teacher.status != "hidden").order_by(Teacher.teacher_no.asc()).all()
    items = [
        {
            "id": teacher.id,
            "teacherNo": teacher.teacher_no,
            "name": teacher.real_name,
            "xileName": teacher.xile_name,
            "tier": teacher.tier.code if teacher.tier else None,
            "tierName": teacher.tier.name if teacher.tier else None,
            "city": teacher.city,
            "district": teacher.district,
            "status": teacher.status,
            "validUntil": _date_text(teacher.valid_until),
            "certifiedAt": _date_text(teacher.first_certified_on),
            "avatarUrl": teacher.avatar_url,
            "phone": teacher.detail.phone if teacher.detail else None,
        }
        for teacher in teachers
    ]
    return {"items": items, "total": len(items)}


@admin_bp.post("/teachers")
@require_admin_token
def create_teacher():
    payload = request.get_json(silent=True) or {}
    name = (payload.get("name") or "").strip()
    if not name:
        return {"error": "name required"}, 400

    tier_code = (payload.get("level") or "L1").strip().upper()
    if not tier_code.startswith("L"):
        tier_code = "L1"
    tier = TeacherTier.query.filter_by(code=tier_code).first()
    if tier is None:
        tier = TeacherTier.query.filter_by(code="L1").first()

    year = date.today().year
    max_no = (
        db.session.query(db.func.max(Teacher.teacher_no))
        .filter(Teacher.teacher_no.like(f"JY{year}%"))
        .scalar()
    )
    if max_no:
        seq = int(max_no[-4:]) + 1
    else:
        seq = 1
    teacher_no = f"JY{year}{seq:04d}"

    valid_until_str = payload.get("expiryDate")
    if valid_until_str:
        try:
            valid_until = date.fromisoformat(valid_until_str.replace(".", "-"))
        except ValueError:
            valid_until = date(year + 3, 12, 31)
    else:
        valid_until = date(year + 3, 12, 31)

    teacher = Teacher(
        teacher_no=teacher_no,
        real_name=name,
        xile_name=payload.get("xileName", "").strip() or None,
        tier_id=tier.id,
        city=payload.get("city", "").strip() or None,
        district=payload.get("district", "").strip() or None,
        status="active",
        first_certified_on=date.today(),
        valid_until=valid_until,
    )
    db.session.add(teacher)
    db.session.flush()

    phone = payload.get("phone", "").strip()
    if phone:
        detail = TeacherDetail(teacher_id=teacher.id, phone=phone)
        db.session.add(detail)

    db.session.commit()

    return {
        "id": teacher.id,
        "teacherNo": teacher.teacher_no,
        "name": teacher.real_name,
        "tier": tier.code,
        "validUntil": _date_text(teacher.valid_until),
    }, 201


@admin_bp.delete("/teachers/<int:teacher_id>")
@require_admin_token
def delete_teacher(teacher_id):
    teacher = db.session.get(Teacher, teacher_id)
    if teacher is None:
        return {"error": "not found"}, 404
    teacher.status = "hidden"
    db.session.add(
        AuditLog(admin_id=1, action="delete_teacher", target_type="teacher", target_id=teacher.id)
    )
    db.session.commit()
    return {"id": teacher.id, "status": "hidden"}


# ─── Studios ─────────────────────────────────────────────────────────────────


@admin_bp.get("/studios")
@require_admin_token
def studio_list():
    studios = Studio.query.order_by(Studio.display_order.desc(), Studio.id.asc()).all()
    items = [
        {
            "id": studio.id,
            "name": studio.name,
            "city": studio.city,
            "district": studio.district,
            "address": studio.address,
            "ownerTeacherName": studio.owner.real_name if studio.owner else None,
            "coverUrl": studio.cover_url,
            "tags": [t.strip() for t in (studio.tags or "").split(",") if t.strip()],
            "intro": studio.intro,
            "openingHours": studio.opening_hours,
            "contactText": studio.contact_text,
            "status": studio.status,
            "displayOrder": studio.display_order,
        }
        for studio in studios
    ]
    return {"items": items, "total": len(items)}


@admin_bp.post("/studios")
@require_admin_token
def create_studio():
    payload = request.get_json(silent=True) or {}
    name = (payload.get("name") or "").strip()
    if not name:
        return {"error": "name required"}, 400

    studio = Studio(
        name=name,
        city=payload.get("city", "").strip() or None,
        district=payload.get("district", "").strip() or None,
        address=payload.get("address", "").strip() or None,
        contact_text=payload.get("contact", "").strip() or None,
        tags=payload.get("tags", "").strip() or None,
        intro=payload.get("intro", "").strip() or None,
        status="open",
    )
    db.session.add(studio)
    db.session.commit()

    return {"id": studio.id, "name": studio.name, "status": studio.status}, 201


@admin_bp.delete("/studios/<int:studio_id>")
@require_admin_token
def delete_studio(studio_id):
    studio = db.session.get(Studio, studio_id)
    if studio is None:
        return {"error": "not found"}, 404
    studio.status = "hidden"
    db.session.commit()
    return {"id": studio.id, "status": "hidden"}


# ─── Reviews ─────────────────────────────────────────────────────────────────


@admin_bp.get("/reviews")
@require_admin_token
def review_queue():
    status = request.args.get("status", "").strip()
    query = AnnualReview.query
    if status and status != "all":
        query = query.filter_by(status=status)
    reviews = query.order_by(AnnualReview.submitted_at.desc(), AnnualReview.id.desc()).all()
    items = [
        {
            "id": review.id,
            "teacherId": review.teacher_id,
            "teacherName": review.teacher.real_name if review.teacher else None,
            "teacherNo": review.teacher.teacher_no if review.teacher else None,
            "xileName": review.teacher.xile_name if review.teacher else None,
            "tier": review.teacher.tier.code if review.teacher and review.teacher.tier else None,
            "city": review.teacher.city if review.teacher else None,
            "avatarUrl": review.teacher.avatar_url if review.teacher else None,
            "reviewYear": review.review_year,
            "status": review.status,
            "submittedAt": _datetime_text(review.submitted_at),
            "reviewedAt": _datetime_text(review.reviewed_at),
            "previousValidUntil": _date_text(review.previous_valid_until),
            "nextValidUntil": _date_text(review.next_valid_until),
            "files": [
                {
                    "id": file.id,
                    "filename": file.file_name,
                    "fileKey": file.file_key,
                    "fileType": file.file_type,
                    "fileSize": file.file_size,
                }
                for file in review.files.order_by("id").all()
            ],
        }
        for review in reviews
    ]
    return {"items": items, "total": len(items)}


@admin_bp.post("/reviews/<int:review_id>/decision")
@require_admin_token
def review_decision(review_id):
    payload = request.get_json(silent=True) or {}
    status = payload.get("status")
    if status not in {"approved", "rejected"}:
        return {"error": "invalid status"}, 400

    review = db.session.get(AnnualReview, review_id)
    if review is None:
        return {"error": "not found"}, 404

    review.status = status
    review.reviewer_comment = payload.get("comment", "")
    review.reviewed_at = datetime.now(timezone.utc).replace(tzinfo=None)

    if status == "approved" and review.next_valid_until and review.teacher:
        review.teacher.valid_until = review.next_valid_until
        if review.teacher.status == "expiring":
            review.teacher.status = "active"

    db.session.commit()

    return {
        "id": review.id,
        "status": review.status,
        "reviewerComment": review.reviewer_comment,
        "reviewedAt": _datetime_text(review.reviewed_at),
    }


# ─── Import ──────────────────────────────────────────────────────────────────


@admin_bp.post("/import/teachers/preview")
@require_admin_token
def import_preview():
    file = request.files.get("file")
    if not file:
        return {"error": "file required"}, 400

    try:
        import openpyxl

        wb = openpyxl.load_workbook(file, read_only=True, data_only=True)
        ws = wb.active
    except Exception:
        return {"error": "invalid excel file"}, 400

    rows = list(ws.iter_rows(min_row=2, values_only=True))
    total_rows = len(rows)
    errors = []
    valid_rows = 0

    for idx, row in enumerate(rows, start=2):
        row_errors = []
        name = str(row[0]).strip() if row[0] else ""
        tier_code = str(row[1]).strip().upper() if row[1] else ""
        cert_date = str(row[2]).strip() if row[2] else ""

        if not name:
            row_errors.append({"rowNumber": idx, "field": "name", "message": "姓名不能为空"})
        if not tier_code or tier_code not in {"L0", "L1", "L2", "L3", "L4", "L5"}:
            row_errors.append({"rowNumber": idx, "field": "level", "message": "等级代码无效（需为L0-L5）"})
        if cert_date:
            try:
                date.fromisoformat(cert_date.replace(".", "-").replace("/", "-"))
            except ValueError:
                row_errors.append({"rowNumber": idx, "field": "certDate", "message": "日期格式无效（需YYYY-MM-DD）"})
        else:
            row_errors.append({"rowNumber": idx, "field": "certDate", "message": "认证日期不能为空"})

        if row_errors:
            errors.extend(row_errors)
        else:
            valid_rows += 1

    batch = ImportBatch(admin_id=1, total_rows=total_rows, success_rows=valid_rows, error_rows=len(errors), status="preview")
    db.session.add(batch)
    db.session.flush()

    for err in errors:
        db.session.add(ImportError(batch_id=batch.id, row_number=err["rowNumber"], field=err["field"], error_message=err["message"]))

    db.session.commit()
    wb.close()

    return {
        "batchId": batch.id,
        "totalRows": total_rows,
        "validRows": valid_rows,
        "errors": errors,
    }


@admin_bp.post("/import/teachers/commit")
@require_admin_token
def import_commit():
    payload = request.get_json(silent=True) or {}
    batch_id = payload.get("batchId")
    if not batch_id:
        return {"error": "batchId required"}, 400

    batch = db.session.get(ImportBatch, batch_id)
    if batch is None or batch.status != "preview":
        return {"error": "invalid batch"}, 400

    file = request.files.get("file")
    if not file:
        return {"error": "file required for commit"}, 400

    try:
        import openpyxl

        wb = openpyxl.load_workbook(file, read_only=True, data_only=True)
        ws = wb.active
    except Exception:
        return {"error": "invalid excel file"}, 400

    rows = list(ws.iter_rows(min_row=2, values_only=True))
    year = date.today().year
    created = 0

    for row in rows:
        name = str(row[0]).strip() if row[0] else ""
        tier_code = str(row[1]).strip().upper() if row[1] else ""
        cert_date_str = str(row[2]).strip() if row[2] else ""
        phone = str(row[3]).strip() if len(row) > 3 and row[3] else ""
        city = str(row[4]).strip() if len(row) > 4 and row[4] else ""

        if not name or tier_code not in {"L0", "L1", "L2", "L3", "L4", "L5"}:
            continue
        try:
            cert_date = date.fromisoformat(cert_date_str.replace(".", "-").replace("/", "-"))
        except ValueError:
            continue

        tier = TeacherTier.query.filter_by(code=tier_code).first()
        if not tier:
            continue

        max_no = (
            db.session.query(db.func.max(Teacher.teacher_no))
            .filter(Teacher.teacher_no.like(f"JY{year}%"))
            .scalar()
        )
        seq = int(max_no[-4:]) + 1 if max_no else 1
        teacher_no = f"JY{year}{seq:04d}"

        cycle = tier.review_cycle_years or 3
        valid_until = date(cert_date.year + cycle, cert_date.month, cert_date.day)

        teacher = Teacher(
            teacher_no=teacher_no,
            real_name=name,
            tier_id=tier.id,
            city=city or None,
            status="active",
            first_certified_on=cert_date,
            valid_until=valid_until,
        )
        db.session.add(teacher)
        db.session.flush()

        if phone:
            db.session.add(TeacherDetail(teacher_id=teacher.id, phone=phone))

        created += 1

    batch.status = "committed"
    batch.success_rows = created
    db.session.commit()
    wb.close()

    return {"batchId": batch.id, "created": created}


# ─── Analytics ───────────────────────────────────────────────────────────────


@admin_bp.get("/analytics")
@require_admin_token
def analytics():
    teachers = Teacher.query.filter(Teacher.status != "hidden").all()
    total = len(teachers)

    tier_counts = {}
    city_counts = {}
    for t in teachers:
        tier_code = t.tier.code if t.tier else "Unknown"
        tier_counts[tier_code] = tier_counts.get(tier_code, 0) + 1
        if t.city:
            city_counts[t.city] = city_counts.get(t.city, 0) + 1

    tier_distribution = [
        {"tier": k, "count": v, "percent": round(v / total * 100, 1) if total else 0}
        for k, v in sorted(tier_counts.items())
    ]
    city_distribution = sorted(city_counts.items(), key=lambda x: -x[1])[:10]
    city_distribution = [
        {"city": k, "count": v, "percent": round(v / total * 100, 1) if total else 0}
        for k, v in city_distribution
    ]

    today = date.today()
    monthly_trend = []
    for i in range(11, -1, -1):
        month_start = date(today.year, today.month, 1)
        month_start = date(
            month_start.year - (i // 12 + (1 if (month_start.month - i % 12) <= 0 else 0)),
            ((month_start.month - i % 12 - 1) % 12) + 1,
            1,
        )
        if month_start.month == 12:
            month_end = date(month_start.year + 1, 1, 1)
        else:
            month_end = date(month_start.year, month_start.month + 1, 1)

        count = AnnualReview.query.filter(
            AnnualReview.status == "approved",
            AnnualReview.reviewed_at >= datetime(month_start.year, month_start.month, month_start.day),
            AnnualReview.reviewed_at < datetime(month_end.year, month_end.month, month_end.day),
        ).count()
        monthly_trend.append({"month": month_start.strftime("%Y-%m"), "count": count})

    active_count = Teacher.query.filter_by(status="active").count()
    total_cities = db.session.query(db.func.count(db.distinct(Teacher.city))).filter(
        Teacher.status != "hidden", Teacher.city.isnot(None)
    ).scalar()

    return {
        "tierDistribution": tier_distribution,
        "cityDistribution": city_distribution,
        "monthlyTrend": monthly_trend,
        "totalTeachers": total,
        "activeRate": round(active_count / total * 100, 1) if total else 0,
        "totalCities": total_cities or 0,
    }


# ─── Settings ────────────────────────────────────────────────────────────────


@admin_bp.get("/settings")
@require_admin_token
def get_settings():
    tiers = TeacherTier.query.order_by(TeacherTier.sort_order.asc()).all()
    tier_settings = [
        {
            "code": t.code,
            "name": t.name,
            "reviewCycleYears": t.review_cycle_years,
            "reviewRequired": t.review_required,
        }
        for t in tiers
    ]

    configs = {c.key: c.value for c in SystemConfig.query.all()}
    return {
        "tiers": tier_settings,
        "features": {
            "qrVerifyEnabled": configs.get("qr_verify_enabled", "true") == "true",
            "certExpiryNotify": configs.get("cert_expiry_notify", "true") == "true",
        },
    }


@admin_bp.put("/settings")
@require_admin_token
def update_settings():
    payload = request.get_json(silent=True) or {}

    tiers_payload = payload.get("tiers")
    if tiers_payload and isinstance(tiers_payload, list):
        for item in tiers_payload:
            tier = TeacherTier.query.filter_by(code=item.get("code")).first()
            if tier and "reviewCycleYears" in item:
                tier.review_cycle_years = item["reviewCycleYears"]

    features = payload.get("features")
    if features and isinstance(features, dict):
        for key, val in features.items():
            db_key = "qr_verify_enabled" if key == "qrVerifyEnabled" else "cert_expiry_notify" if key == "certExpiryNotify" else None
            if db_key:
                config = db.session.get(SystemConfig, db_key)
                if config:
                    config.value = "true" if val else "false"
                else:
                    db.session.add(SystemConfig(key=db_key, value="true" if val else "false"))

    db.session.commit()
    return {"ok": True}


# ─── Permissions ─────────────────────────────────────────────────────────────


@admin_bp.get("/permissions")
@require_admin_token
def get_permissions():
    admins = AdminUser.query.order_by(AdminUser.id.asc()).all()
    items = [
        {
            "id": a.id,
            "username": a.username,
            "role": a.role,
            "createdAt": _datetime_text(a.created_at),
        }
        for a in admins
    ]
    return {"items": items, "total": len(items)}


@admin_bp.post("/permissions/invite")
@require_admin_token
def invite_admin():
    payload = request.get_json(silent=True) or {}
    username = (payload.get("username") or "").strip()
    if not username:
        return {"error": "username required"}, 400

    existing = AdminUser.query.filter_by(username=username).first()
    if existing:
        return {"error": "username already exists"}, 409

    from werkzeug.security import generate_password_hash

    password = payload.get("password") or "changeme123"
    role = payload.get("role") or "admin"

    admin = AdminUser(
        username=username,
        password_hash=generate_password_hash(password),
        role=role,
    )
    db.session.add(admin)
    db.session.commit()

    return {"id": admin.id, "username": admin.username, "role": admin.role}, 201
